r"""Batch-process every PDF in a folder: extract 姓名/电话/邮箱/日期 via DeepSeek,
aggregate into a pandas DataFrame, and export to employees.xlsx.

Usage:
    python batch_process.py                       # 默认处理 .\data，输出 .\employees.xlsx
    python batch_process.py --input data --output employees.xlsx
"""

from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

try:
    import pandas as pd
except ImportError:
    sys.exit("缺少 pandas，请先执行：python -m pip install pandas openpyxl")

try:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("缺少 openpyxl，请先执行：python -m pip install pandas openpyxl")

import info_extractor as ie

HERE = Path(__file__).resolve().parent
COLUMNS = ["文件名", "姓名", "电话", "邮箱", "日期", "状态"]
COLUMN_WIDTHS = {"文件名": 24, "姓名": 12, "电话": 34, "邮箱": 32, "日期": 26, "状态": 40}


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="批量提取 data 目录下所有 PDF 的信息并导出 Excel")
    parser.add_argument("--input", default=str(HERE / "data"), help="PDF 所在目录，默认 ./data")
    parser.add_argument("--output", default=str(HERE / "employees.xlsx"), help="输出 Excel 路径")
    parser.add_argument("--json", default=None, help="可选：同时导出一份汇总 JSON")
    parser.add_argument("--api-key", default=None, help="DeepSeek API Key，默认读取 DEEPSEEK_API_KEY")
    parser.add_argument("--model", default=ie.MODEL, help=f"模型名，默认 {ie.MODEL}")
    parser.add_argument("--delay", type=float, default=1.0, help="每个文件之间的间隔秒数，默认 1（防接口限流）")
    parser.add_argument("--keep-raw", action="store_true", help="额外输出一列 _raw（模型原始 JSON）")
    return parser.parse_args(argv[1:])


def collect_pdfs(input_dir: Path) -> list[Path]:
    """Return all PDFs in input_dir, sorted by name."""
    return sorted((p for p in input_dir.iterdir() if p.is_file() and p.suffix.lower() == ".pdf"),
                  key=lambda p: p.name)


def process_one(pdf_path: Path, api_key: str, model: str, index: int, total: int) -> dict:
    """Extract one PDF; never raises — returns a row dict with 状态 describing the outcome.

    缺失的字段一律填 None（导出 Excel / JSON 时为 null），空字符串也统一转成 None。
    """
    row: dict = {"文件名": pdf_path.name, "姓名": None, "电话": None, "邮箱": None, "日期": None, "状态": "ok"}
    print(f"[{index}/{total}] {pdf_path.name} ... ", end="", flush=True)
    try:
        text = ie.extract_pdf_text(pdf_path)
        if not text:
            row["状态"] = "跳过：PDF 无可提取文本（可能是扫描件）"
            print("跳过（无文本）")
            return row
        data = ie.extract_fields(text, api_key, model)
        for field in ie.FIELDS:
            value = data.get(field) or ""
            row[field] = value if str(value).strip() else None
        if args_keep_raw:
            row["_raw"] = data.get("_raw", "")
        print("完成")
    except Exception as exc:
        row["状态"] = f"{exc.__class__.__name__}: {exc}"
        print(f"失败（{exc.__class__.__name__}）")
    return row


args_keep_raw = False


def to_json_records(df: pd.DataFrame) -> list[dict]:
    """DataFrame -> list of dicts where missing values become real JSON null (not NaN)."""
    def clean(value):
        if value is None:
            return None
        if isinstance(value, float) and value != value:  # NaN
            return None
        return value

    return [{str(k): clean(v) for k, v in record.items()} for record in df.to_dict(orient="records")]


def build_dataframe(rows: list[dict], keep_raw: bool) -> pd.DataFrame:
    columns = list(COLUMNS) + (["_raw"] if keep_raw else [])
    return pd.DataFrame(rows, columns=columns)


def export_excel(df: pd.DataFrame, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="employees")
        sheet = writer.sheets["employees"]

        header_fill = PatternFill("solid", fgColor="1F4E79")
        for col_idx, column in enumerate(df.columns, start=1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            sheet.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS.get(column, 18)

        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions


def main(argv: list[str]) -> int:
    global args_keep_raw
    args = parse_args(argv)
    args_keep_raw = args.keep_raw

    input_dir = Path(args.input).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()

    if not input_dir.is_dir():
        print(f"[错误] 输入目录不存在：{input_dir}")
        print(f"请创建该目录并放入 PDF，例如：\n  mkdir \"{input_dir}\"")
        print(f"  .venv\\Scripts\\python.exe make_test_pdfs.py \"{input_dir}\"   :: 生成测试 PDF")
        return 1

    pdfs = collect_pdfs(input_dir)
    if not pdfs:
        print(f"[错误] 目录里没有 PDF 文件：{input_dir}")
        print(f"请放入 PDF，或先生成测试文件：\n  .venv\\Scripts\\python.exe make_test_pdfs.py \"{input_dir}\"")
        return 1

    api_key = ie.resolve_api_key(args.api_key, fallback_dir=HERE)
    if not api_key:
        print(
            "[错误] 未找到 DeepSeek API Key。\n"
            '  cmd:        set DEEPSEEK_API_KEY=sk-xxxx\n'
            '  PowerShell: $env:DEEPSEEK_API_KEY="sk-xxxx"\n'
            f"  或在 {HERE}\\.env 写入：DEEPSEEK_API_KEY=sk-xxxx"
        )
        return 1

    print(f"输入目录：{input_dir}")
    print(f"发现 {len(pdfs)} 个 PDF，模型：{args.model}\n")

    rows: list[dict] = []
    for index, pdf_path in enumerate(pdfs, start=1):
        rows.append(process_one(pdf_path, api_key, args.model, index, len(pdfs)))
        if args.delay and index < len(pdfs):
            time.sleep(args.delay)

    df = build_dataframe(rows, args.keep_raw)
    export_excel(df, output_path)

    if args.json:
        json_path = Path(args.json).expanduser().resolve()
        json_path.write_text(
            json.dumps(to_json_records(df), ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print(f"\n汇总 JSON 已写入：{json_path}")

    ok = int((df["状态"] == "ok").sum())
    print("\n===== 汇总结果 =====")
    print(df.to_string(index=False, max_colwidth=28))
    print(f"\n成功 {ok}/{len(df)} 个文件")
    print(f"Excel 已生成：{output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
